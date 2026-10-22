#!/usr/bin/env python3
# MIT License. Shared helpers for the docx skill scripts.
"""Shared helpers: paragraph iteration and run-preserving text replacement."""
from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run


class _AliasDedup:
    """Identity dedup over underlying XML nodes for ONE traversal invocation.

    Merged table cells (horizontal ``gridSpan``, vertical ``vMerge`` and
    rectangular merges) expose the same ``w:tc`` subtree once per grid slot,
    and linked header/footer parts expose the same part element once per
    linking section, so the same underlying ``w:p`` node can be reached many
    times in one traversal.  Keys are ``id()`` of the lxml proxy objects;
    lxml returns the same proxy for the same underlying node while a
    reference to it is alive, and this class holds a strong reference to
    every recorded node for the whole traversal, so no proxy identity can be
    garbage-collected and recycled for a different node.  Identity, never
    text or XML serialization, decides duplication: distinct nodes with
    identical content stay distinct.
    """

    def __init__(self):
        self._ids = set()
        self._keepalive = []

    def node_is_new(self, element) -> bool:
        if id(element) in self._ids:
            return False
        self._ids.add(id(element))
        self._keepalive.append(element)
        return True

    def paragraph_is_new(self, para) -> bool:
        return self.node_is_new(para._p)

    def cell_is_new(self, cell) -> bool:
        return self.node_is_new(cell._tc)


def iter_all_paragraphs(doc, include_headers_footers: bool = True):
    """Yield every paragraph in body, tables (recursively), headers, footers.

    Each underlying ``w:p`` XML node is yielded exactly once per invocation,
    in the original first-encounter order (paragraphs first, then tables,
    per container; body first, then each section's header/footer parts).
    Merged-cell aliases and shared/linked header/footer parts no longer
    rescan or re-yield a node already yielded.  ``include_headers_footers``
    keeps its original meaning.
    """
    dedup = _AliasDedup()
    yield from _iter_container(doc, dedup)
    if include_headers_footers:
        seen_parts = _AliasDedup()
        for section in doc.sections:
            for part in (
                section.header, section.footer,
                section.first_page_header, section.first_page_footer,
                section.even_page_header, section.even_page_footer,
            ):
                if part is not None and seen_parts.node_is_new(part._element):
                    yield from _iter_container(part, dedup)


def _iter_container(container, dedup=None):
    if dedup is None:
        dedup = _AliasDedup()
    for para in container.paragraphs:
        if dedup.paragraph_is_new(para):
            yield para
    for table in container.tables:
        yield from _iter_table(table, dedup)


def _iter_table(table, dedup=None):
    if dedup is None:
        dedup = _AliasDedup()
    for row in table.rows:
        for cell in row.cells:
            if not dedup.cell_is_new(cell):
                continue  # merged-cell alias of an already visited w:tc
            for para in cell.paragraphs:
                if dedup.paragraph_is_new(para):
                    yield para
            for nested in cell.tables:
                yield from _iter_table(nested, dedup)


def iter_part_roots(doc):
    """Yield the XML root of the body plus every header/footer part."""
    yield doc.element.body
    seen = set()
    for section in doc.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if part is not None and id(part._element) not in seen:
                seen.add(id(part._element))
                yield part._element


def _direct_run_segments(para):
    """Group ``para.runs`` indices into maximal segments of consecutive
    ``w:r`` children of the paragraph element.

    ``Paragraph.runs`` covers only direct ``w:r`` children; runs nested in
    ``w:hyperlink`` or other wrapper elements are excluded, so concatenating
    every run text would invent adjacency that does not exist in the real
    paragraph text.  Every non-``w:r`` child (``w:hyperlink``, bookmarks,
    ``w:sdt``, ``w:pPr``, ...) is a hard barrier: matches never cross it and the
    barrier XML stays byte-exact.  Hyperlink-internal text is matched
    separately, wholly inside each direct ``w:hyperlink`` container (see
    ``replace_in_paragraph``); text inside other wrappers stays explicitly
    unsupported instead of being silently corrupted.
    """
    segments = []
    current = []
    index = 0
    for child in para._p.iterchildren():
        if child.tag == qn('w:r'):
            current.append(index)
            index += 1
        elif current:
            segments.append(current)
            current = []
    if current:
        segments.append(current)
    return segments


_RUN_TEXT_TAGS = frozenset(qn('w:' + name) for name in (
    't', 'tab', 'br', 'cr', 'noBreakHyphen', 'ptab',
))


def _rewrite_text_child(child, text):
    """Change one text-bearing child, retaining its original XML position.

    Use python-docx's text encoder on a detached, empty run so inserted tabs
    and CR/LF have the same XML/text semantics as Run.text.  Never clear the
    real run.  Ordinary w:t edits retain that element and its attributes;
    generated w:t elements retain its attributes as well as xml:space when
    the new text requires it.  Non-text siblings are never moved or copied.
    """
    if text == str(child):
        return
    if child.tag == qn('w:t') and not any(c in text for c in '\t\r\n'):
        child.text = text or None
        if len(text.strip()) < len(text):
            child.set(qn('xml:space'), 'preserve')
        return
    holder = OxmlElement('w:r')
    holder.text = text
    generated = list(holder)
    if not generated:
        generated = [OxmlElement('w:t')]
    for element in generated:
        if child.tag == element.tag == qn('w:t'):
            element.attrib.update(child.attrib)
            value = element.text or ''
            if len(value.strip()) < len(value):
                element.set(qn('xml:space'), 'preserve')
        child.addprevious(element)
    generated[-1].tail = child.tail
    child.getparent().remove(child)


def _replace_run_text_preserving_xml(run, edits):
    """Apply sorted original-local edits once across this run's text children.

    Opaque children (including fields, drawings, rPr and zero-text breaks)
    remain the same nodes in the same run and original sibling order.
    Within each edited run, original character boundaries keep
    their relative replacement offset, clamped to its length; any expansion
    remainder belongs to the last overlapped original character.  Equal text
    therefore keeps internal opaque boundaries unchanged.  A contraction or
    deletion can collapse boundaries but never reverses, moves or deletes
    their XML nodes.  The caller distributes replacement text across runs
    when an opaque boundary is crossed; otherwise the existing first-run
    formatting policy applies.  Unaffected text children receive zero writes.
    """
    offset = 0
    edit_index = 0
    for child in list(run._r):
        if child.tag not in _RUN_TEXT_TAGS:
            continue
        text = str(child)
        end = offset + len(text)
        if not text:  # e.g. page/column br: opaque to Run.text matching
            continue
        while edit_index < len(edits) and edits[edit_index][1] <= offset:
            edit_index += 1
        index = edit_index
        parts = []
        cursor = 0
        while index < len(edits) and edits[index][0] < end:
            start_cut, end_cut, replacement = edits[index]
            parts.append(text[cursor:max(start_cut - offset, 0)])
            if replacement:
                # Preserve each original internal character boundary when
                # possible.  Only the final overlapped character receives
                # expansion beyond the original matched extent of this run.
                replacement_start = max(start_cut, offset) - start_cut
                replacement_end = (len(replacement) if end_cut <= end else
                                   end - start_cut)
                parts.append(replacement[replacement_start:replacement_end])
            cursor = min(end_cut, end) - offset
            if end_cut > end:
                break
            index += 1
        if index != edit_index or cursor:
            parts.append(text[cursor:])
            _rewrite_text_child(child, ''.join(parts))
        edit_index = index
        offset = end


def _hyperlink_run_groups(para):
    """Yield ``(runs, segments)`` for each direct ``w:hyperlink`` child.

    Only hyperlinks that are DIRECT children of the paragraph element are
    eligible; hyperlinks nested inside other wrappers (``w:sdt``,
    ``w:smartTag``, ...) stay explicitly unsupported and are never entered.
    Within one hyperlink only its direct ``w:r`` children participate,
    grouped into maximal consecutive-run segments exactly like paragraph
    direct runs; every non-``w:r`` child is a hard barrier.  Each hyperlink
    is its own group, so matches are invented neither across two hyperlinks
    nor between direct runs and hyperlink runs.  Hyperlink attributes,
    relationship ids/targets and run ``rPr`` are never modified: only
    text-bearing children of positively overlapped runs change.
    """
    for child in para._p.iterchildren():
        if child.tag != qn('w:hyperlink'):
            continue
        runs = []
        segments = []
        current = []
        for element in child.iterchildren():
            if element.tag == qn('w:r'):
                current.append(len(runs))
                runs.append(Run(element, para))
            elif current:
                segments.append(current)
                current = []
        if current:
            segments.append(current)
        if segments:
            yield runs, segments


def _replace_in_run_segments(runs, segments, old, new):
    """Original matching/edit logic shared by direct runs and each eligible
    hyperlink container.  Callers guarantee a non-empty ``old``."""
    texts = [r.text for r in runs]
    edits = [[] for _ in runs]
    total = 0
    for segment in segments:
        seg_texts = [texts[i] for i in segment]
        offsets = []  # original start offset of each run in the segment text
        pos = 0
        for t in seg_texts:
            offsets.append(pos)
            pos += len(t)
        full = "".join(seg_texts)
        starts = []
        pos = 0
        while True:
            start = full.find(old, pos)
            if start < 0:
                break
            starts.append(start)
            pos = start + len(old)  # non-overlapping
        # Global visible-text offsets of opaque children, including markers
        # in empty runs.  Runs/properties themselves are not boundaries.
        opaque_offsets = []
        for i, run_index in enumerate(segment):
            child_offset = offsets[i]
            for child in runs[run_index]._r:
                if child.tag in _RUN_TEXT_TAGS:
                    value = str(child)
                    if value:
                        child_offset += len(value)
                    elif child.tag == qn('w:br'):
                        opaque_offsets.append(child_offset)  # page/column break
                elif child.tag != qn('w:rPr'):
                    opaque_offsets.append(child_offset)
        # Collect edits per original run as (cut_start, cut_end,
        # replacement|None) in original run-local offsets.  Matches and run
        # offsets are both sorted, so a single forward cursor locates the
        # overlapping runs of each match.  A monotone opaque-boundary cursor
        # decides whether replacement must retain global character slots.
        n = len(segment)
        lo = 0
        boundary_index = 0
        for start in starts:
            end = start + len(old)
            while (boundary_index < len(opaque_offsets) and
                   opaque_offsets[boundary_index] <= start):
                boundary_index += 1
            preserve_anchors = (boundary_index < len(opaque_offsets) and
                                opaque_offsets[boundary_index] < end)
            while lo < n and offsets[lo] + len(seg_texts[lo]) <= start:
                lo += 1
            first = True
            i = lo
            while i < n and offsets[i] < end:
                r_end = offsets[i] + len(seg_texts[i])
                cs = max(start, offsets[i]) - offsets[i]
                ce = min(end, r_end) - offsets[i]
                if ce > cs:  # positive overlap only; empty/boundary runs untouched
                    if preserve_anchors:
                        relative_start = max(start, offsets[i]) - start
                        relative_end = (len(new) if r_end >= end else
                                        r_end - start)
                        replacement = new[relative_start:relative_end]
                    else:
                        replacement = new if first else None
                    edits[segment[i]].append((cs, ce, replacement))
                    first = False
                i += 1
        total += len(starts)
    # Reconstruct each affected run once; retain opaque nodes and unchanged
    # text children in place.  Untouched/empty runs receive zero writes.
    for i, run_edits in enumerate(edits):
        if not run_edits:
            continue
        _replace_run_text_preserving_xml(runs[i], run_edits)
    return total


def replace_in_paragraph(para, old: str, new: str) -> int:
    """Replace `old` with `new` in a paragraph, preserving run formatting.

    Eligible text lives in two kinds of containers that never share a match:
    maximal segments of consecutive direct ``w:r`` children of the paragraph
    (original behaviour, unchanged), and -- new in this candidate -- each
    DIRECT ``w:hyperlink`` child of the paragraph, matched wholly inside its
    own direct ``w:r`` children under the identical non-overlapping
    original-text policy.  Hyperlink attributes, relationship ids/targets,
    ``rPr`` and all opaque nodes are preserved; matches never cross a
    container boundary in either direction, never enter hyperlinks nested in
    other wrappers (``w:sdt``, ``w:smartTag``, ...), and never rescan
    inserted replacement text, so the function terminates even when ``new``
    contains ``old`` or ``old == new``.

    Formatting policy inside each container is exactly the pre-existing one:
    ordinary matches inherit the first positively overlapped run's
    formatting; matches crossing an opaque XML boundary keep global
    character anchors.  Returns number of replacements made.  No field
    recalculation or Office rendering is claimed.
    """
    if not old:
        return 0
    total = _replace_in_run_segments(para.runs, _direct_run_segments(para),
                                     old, new)
    for hl_runs, hl_segments in _hyperlink_run_groups(para):
        total += _replace_in_run_segments(hl_runs, hl_segments, old, new)
    return total
