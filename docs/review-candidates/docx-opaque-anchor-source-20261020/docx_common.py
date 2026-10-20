#!/usr/bin/env python3
# MIT License. Shared helpers for the docx skill scripts.
"""Shared helpers: paragraph iteration and run-preserving text replacement."""
from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def iter_all_paragraphs(doc, include_headers_footers: bool = True):
    """Yield every paragraph in body, tables (recursively), headers, footers."""
    yield from _iter_container(doc)
    if include_headers_footers:
        for section in doc.sections:
            for part in (
                section.header, section.footer,
                section.first_page_header, section.first_page_footer,
                section.even_page_header, section.even_page_footer,
            ):
                if part is not None:
                    yield from _iter_container(part)


def _iter_container(container):
    for para in container.paragraphs:
        yield para
    for table in container.tables:
        yield from _iter_table(table)


def _iter_table(table):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                yield para
            for nested in cell.tables:
                yield from _iter_table(nested)


def iter_part_roots(doc):
    """Yield the XML root of the body plus every header/footer part."""
    yield doc.element.body
    seen = set()
    for section in doc.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if part is not None and id(part._element) not in seen:
                seen.add(id(part._element))
                yield part._element


def _direct_run_segments(para):
    """Group ``para.runs`` indices into maximal segments of consecutive
    ``w:r`` children of the paragraph element.

    ``Paragraph.runs`` covers only direct ``w:r`` children; runs nested in
    ``w:hyperlink`` or other wrapper elements are excluded, so concatenating
    every run text would invent adjacency that does not exist in the real
    paragraph text.  Every non-``w:r`` child (``w:hyperlink``, bookmarks,
    ``w:sdt``, ``w:pPr``, ...) is a hard barrier: matches never cross it, the
    barrier XML stays byte-exact, and hyperlink-internal or cross-container
    text remains explicitly unsupported instead of being silently corrupted.
    """
    segments = []
    current = []
    index = 0
    for child in para._p.iterchildren():
        if child.tag == qn('w:r'):
            current.append(index)
            index += 1
        elif current:
            segments.append(current)
            current = []
    if current:
        segments.append(current)
    return segments


_RUN_TEXT_TAGS = frozenset(qn('w:' + name) for name in (
    't', 'tab', 'br', 'cr', 'noBreakHyphen', 'ptab',
))


def _rewrite_text_child(child, text):
    """Change one text-bearing child, retaining its original XML position.

    Use python-docx's text encoder on a detached, empty run so inserted tabs
    and CR/LF have the same XML/text semantics as Run.text.  Never clear the
    real run.  Ordinary w:t edits retain that element and its attributes;
    generated w:t elements retain its attributes as well as xml:space when
    the new text requires it.  Non-text siblings are never moved or copied.
    """
    if text == str(child):
        return
    if child.tag == qn('w:t') and not any(c in text for c in '\t\r\n'):
        child.text = text or None
        if len(text.strip()) < len(text):
            child.set(qn('xml:space'), 'preserve')
        return
    holder = OxmlElement('w:r')
    holder.text = text
    generated = list(holder)
    if not generated:
        generated = [OxmlElement('w:t')]
    for element in generated:
        if child.tag == element.tag == qn('w:t'):
            element.attrib.update(child.attrib)
            value = element.text or ''
            if len(value.strip()) < len(value):
                element.set(qn('xml:space'), 'preserve')
        child.addprevious(element)
    generated[-1].tail = child.tail
    child.getparent().remove(child)


def _replace_run_text_preserving_xml(run, edits):
    """Apply sorted original-local edits once across this run's text children.

    Opaque children (including fields, drawings, rPr and zero-text breaks)
    remain the same nodes in the same run and original sibling order.
    Within each edited run, original character boundaries keep
    their relative replacement offset, clamped to its length; any expansion
    remainder belongs to the last overlapped original character.  Equal text
    therefore keeps internal opaque boundaries unchanged.  A contraction or
    deletion can collapse boundaries but never reverses, moves or deletes
    their XML nodes.  The caller distributes replacement text across runs
    when an opaque boundary is crossed; otherwise the existing first-run
    formatting policy applies.  Unaffected text children receive zero writes.
    """
    offset = 0
    edit_index = 0
    for child in list(run._r):
        if child.tag not in _RUN_TEXT_TAGS:
            continue
        text = str(child)
        end = offset + len(text)
        if not text:  # e.g. page/column br: opaque to Run.text matching
            continue
        while edit_index < len(edits) and edits[edit_index][1] <= offset:
            edit_index += 1
        index = edit_index
        parts = []
        cursor = 0
        while index < len(edits) and edits[index][0] < end:
            start_cut, end_cut, replacement = edits[index]
            parts.append(text[cursor:max(start_cut - offset, 0)])
            if replacement:
                # Preserve each original internal character boundary when
                # possible.  Only the final overlapped character receives
                # expansion beyond the original matched extent of this run.
                replacement_start = max(start_cut, offset) - start_cut
                replacement_end = (len(replacement) if end_cut <= end else
                                   end - start_cut)
                parts.append(replacement[replacement_start:replacement_end])
            cursor = min(end_cut, end) - offset
            if end_cut > end:
                break
            index += 1
        if index != edit_index or cursor:
            parts.append(text[cursor:])
            _rewrite_text_child(child, ''.join(parts))
        edit_index = index
        offset = end


def replace_in_paragraph(para, old: str, new: str) -> int:
    """Replace `old` with `new` in a paragraph, preserving run formatting.

    Matches are the non-overlapping occurrences of `old` in the paragraph's
    ORIGINAL text, scanned left to right within each maximal segment of
    consecutive direct runs; non-run XML children of the paragraph
    (hyperlinks, bookmarks, ...) are barriers that matches never cross.  Each
    match is processed exactly once and inserted replacement text is never
    rescanned, so the function terminates even when `new` contains `old` or
    `old == new`.

    Formatting policy: ordinary matches inherit the first positively
    overlapped run's formatting.  If a match crosses an opaque XML boundary
    (fields, drawings, etc., including in empty runs), preserving its global
    anchor takes priority: replacement characters occupy their original
    character slots across runs, clamped to replacement length, and expansion
    remainder belongs to the last matched character.  Original run styles
    therefore apply on each side of that boundary.  Equal opaque-spanning
    replacements keep XML unchanged; shorter replacements can collapse
    boundaries without moving, deleting or reversing their XML nodes.  rPr
    is formatting, not an opaque boundary for this policy.  Returns number
    of replacements made.  No field recalculation or Office rendering is
    claimed.

    Implementation: original run offsets are mapped once per segment, matches
    are found on the original segment text, edits are collected per original
    run only for positive-overlap ranges (empty or boundary-touching runs
    receive no edit and no assignment, so their XML -- e.g. a field marker in
    an empty run -- stays byte-identical).  One forward reconstruction per
    affected run changes only positively overlapped text-bearing children;
    Run.text is never assigned on the real run.  Work is proportional to
    input + output + XML children + runs + matches; runs are not rescanned
    per match and no per-match full-string rebuilding occurs.
    """
    if not old:
        return 0
    runs = para.runs
    texts = [r.text for r in runs]
    edits = [[] for _ in runs]
    total = 0
    for segment in _direct_run_segments(para):
        seg_texts = [texts[i] for i in segment]
        offsets = []  # original start offset of each run in the segment text
        pos = 0
        for t in seg_texts:
            offsets.append(pos)
            pos += len(t)
        full = "".join(seg_texts)
        starts = []
        pos = 0
        while True:
            start = full.find(old, pos)
            if start < 0:
                break
            starts.append(start)
            pos = start + len(old)  # non-overlapping
        # Global visible-text offsets of opaque children, including markers
        # in empty runs.  Runs/properties themselves are not boundaries.
        opaque_offsets = []
        for i, run_index in enumerate(segment):
            child_offset = offsets[i]
            for child in runs[run_index]._r:
                if child.tag in _RUN_TEXT_TAGS:
                    value = str(child)
                    if value:
                        child_offset += len(value)
                    elif child.tag == qn('w:br'):
                        opaque_offsets.append(child_offset)  # page/column break
                elif child.tag != qn('w:rPr'):
                    opaque_offsets.append(child_offset)
        # Collect edits per original run as (cut_start, cut_end,
        # replacement|None) in original run-local offsets.  Matches and run
        # offsets are both sorted, so a single forward cursor locates the
        # overlapping runs of each match.  A monotone opaque-boundary cursor
        # decides whether replacement must retain global character slots.
        n = len(segment)
        lo = 0
        boundary_index = 0
        for start in starts:
            end = start + len(old)
            while (boundary_index < len(opaque_offsets) and
                   opaque_offsets[boundary_index] <= start):
                boundary_index += 1
            preserve_anchors = (boundary_index < len(opaque_offsets) and
                                opaque_offsets[boundary_index] < end)
            while lo < n and offsets[lo] + len(seg_texts[lo]) <= start:
                lo += 1
            first = True
            i = lo
            while i < n and offsets[i] < end:
                r_end = offsets[i] + len(seg_texts[i])
                cs = max(start, offsets[i]) - offsets[i]
                ce = min(end, r_end) - offsets[i]
                if ce > cs:  # positive overlap only; empty/boundary runs untouched
                    if preserve_anchors:
                        relative_start = max(start, offsets[i]) - start
                        relative_end = (len(new) if r_end >= end else
                                        r_end - start)
                        replacement = new[relative_start:relative_end]
                    else:
                        replacement = new if first else None
                    edits[segment[i]].append((cs, ce, replacement))
                    first = False
                i += 1
        total += len(starts)
    # Reconstruct each affected run once; retain opaque nodes and unchanged
    # text children in place.  Untouched/empty runs receive zero writes.
    for i, run_edits in enumerate(edits):
        if not run_edits:
            continue
        _replace_run_text_preserving_xml(runs[i], run_edits)
    return total
