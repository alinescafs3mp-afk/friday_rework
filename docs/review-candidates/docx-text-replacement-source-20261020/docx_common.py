#!/usr/bin/env python3
# MIT License. Shared helpers for the docx skill scripts.
"""Shared helpers: paragraph iteration and run-preserving text replacement."""
from __future__ import annotations

from docx.oxml.ns import qn


def iter_all_paragraphs(doc, include_headers_footers: bool = True):
    """Yield every paragraph in body, tables (recursively), headers, footers."""
    yield from _iter_container(doc)
    if include_headers_footers:
        for section in doc.sections:
            for part in (
                section.header, section.footer,
                section.first_page_header, section.first_page_footer,
                section.even_page_header, section.even_page_footer,
            ):
                if part is not None:
                    yield from _iter_container(part)


def _iter_container(container):
    for para in container.paragraphs:
        yield para
    for table in container.tables:
        yield from _iter_table(table)


def _iter_table(table):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                yield para
            for nested in cell.tables:
                yield from _iter_table(nested)


def iter_part_roots(doc):
    """Yield the XML root of the body plus every header/footer part."""
    yield doc.element.body
    seen = set()
    for section in doc.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if part is not None and id(part._element) not in seen:
                seen.add(id(part._element))
                yield part._element


def _direct_run_segments(para):
    """Group ``para.runs`` indices into maximal segments of consecutive
    ``w:r`` children of the paragraph element.

    ``Paragraph.runs`` covers only direct ``w:r`` children; runs nested in
    ``w:hyperlink`` or other wrapper elements are excluded, so concatenating
    every run text would invent adjacency that does not exist in the real
    paragraph text.  Every non-``w:r`` child (``w:hyperlink``, bookmarks,
    ``w:sdt``, ``w:pPr``, ...) is a hard barrier: matches never cross it, the
    barrier XML stays byte-exact, and hyperlink-internal or cross-container
    text remains explicitly unsupported instead of being silently corrupted.
    """
    segments = []
    current = []
    index = 0
    for child in para._p.iterchildren():
        if child.tag == qn('w:r'):
            current.append(index)
            index += 1
        elif current:
            segments.append(current)
            current = []
    if current:
        segments.append(current)
    return segments


def replace_in_paragraph(para, old: str, new: str) -> int:
    """Replace `old` with `new` in a paragraph, preserving run formatting.

    Matches are the non-overlapping occurrences of `old` in the paragraph's
    ORIGINAL text, scanned left to right within each maximal segment of
    consecutive direct runs; non-run XML children of the paragraph
    (hyperlinks, bookmarks, ...) are barriers that matches never cross.  Each
    match is processed exactly once and inserted replacement text is never
    rescanned, so the function terminates even when `new` contains `old` or
    `old == new`.

    Formatting policy (unchanged): an occurrence fully contained in a single
    run keeps that run's formatting; an occurrence spanning multiple runs
    collapses the matched runs, the replacement inheriting the formatting of
    the run where the match starts.  Returns number of replacements made.

    Implementation: original run offsets are mapped once per segment, matches
    are found on the original segment text, edits are collected per original
    run only for positive-overlap ranges (empty or boundary-touching runs
    receive no edit and no assignment, so their XML -- e.g. a field marker in
    an empty run -- stays byte-identical), and each affected Run.text is
    reconstructed and assigned exactly once.  Work is proportional to input +
    output + runs + matches; runs are not rescanned per match and no
    per-match full-string rebuilding occurs.
    """
    if not old:
        return 0
    runs = para.runs
    texts = [r.text for r in runs]
    edits = [[] for _ in runs]
    total = 0
    for segment in _direct_run_segments(para):
        seg_texts = [texts[i] for i in segment]
        offsets = []  # original start offset of each run in the segment text
        pos = 0
        for t in seg_texts:
            offsets.append(pos)
            pos += len(t)
        full = "".join(seg_texts)
        starts = []
        pos = 0
        while True:
            start = full.find(old, pos)
            if start < 0:
                break
            starts.append(start)
            pos = start + len(old)  # non-overlapping
        # Collect edits per original run as (cut_start, cut_end,
        # replacement|None) in original run-local offsets.  Matches and run
        # offsets are both sorted, so a single forward cursor locates the
        # overlapping runs of each match; the replacement text is inserted
        # only in the first run with positive original-text overlap.
        n = len(segment)
        lo = 0
        for start in starts:
            end = start + len(old)
            while lo < n and offsets[lo] + len(seg_texts[lo]) <= start:
                lo += 1
            first = True
            i = lo
            while i < n and offsets[i] < end:
                r_end = offsets[i] + len(seg_texts[i])
                cs = max(start, offsets[i]) - offsets[i]
                ce = min(end, r_end) - offsets[i]
                if ce > cs:  # positive overlap only; empty/boundary runs untouched
                    edits[segment[i]].append((cs, ce, new if first else None))
                    first = False
                i += 1
        total += len(starts)
    # Reconstruct each affected run from its original text exactly once;
    # untouched runs are never assigned, keeping their XML byte-identical.
    for i, run_edits in enumerate(edits):
        if not run_edits:
            continue
        t = texts[i]
        parts = []
        cursor = 0
        for cs, ce, repl in run_edits:
            parts.append(t[cursor:cs])
            if repl:
                parts.append(repl)
            cursor = ce
        parts.append(t[cursor:])
        runs[i].text = "".join(parts)
    return total
